from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest
from streamlit.runtime.memory_media_file_storage import MemoryMediaFileStorage
from streamlit.testing.v1 import app_test

from result_chart_helpers import chart_event, table_event
from test_grouping import completed, open_results
from src.category_summary import OTHER_SENTIMENT
from src.response_table import save_table_sentiment
from src.result_downloads import classification_frame, statistics_frames, xlsx_download
from src.result_groups import initial_layout, load_code_group
from src.models import CodingResult, Issue
from src.results import result_tables


def test_voc_download_preserves_multiple_classes_separate_ids_and_unclassified_rows(completed):
    store, run_id, _ = completed
    originals, issues = result_tables(store, run_id)
    before = originals.copy(deep=True), issues.copy(deep=True)
    frame = classification_frame(originals, issues)
    assert list(frame.columns) == ["VOC ID", "VOC 원문", "대분류", "세부분류", "감성"]
    assert frame["VOC ID"].tolist() == ["V0001", "V0001", "V0002", "V0003", "V0004", "V0005"]
    assert frame.loc[frame["VOC ID"].eq("V0001"), "세부분류"].tolist() == ["배송 속도", "오배송"]
    assert frame.loc[frame["VOC ID"].eq("V0005"), ["대분류", "세부분류", "감성"]].iloc[0].tolist() == ["", "", ""]
    pd.testing.assert_frame_equal(originals, before[0])
    pd.testing.assert_frame_equal(issues, before[1])


def test_statistics_match_user_example_and_keep_sentiment_denominators_separate(completed):
    store, run_id, codes = completed
    sheets = statistics_frames(*result_tables(store, run_id), codes)
    assert list(sheets) == ["대분류", "세부분류"]
    assert sheets["대분류"].iloc[0].tolist() == ["배송", 4, 80, 25, 50, 25, 0]
    assert sheets["세부분류"].iloc[0].tolist() == ["배송", "오배송", 3, 60, 0, 100, 0, 0]
    assert sheets["세부분류"].iloc[1].tolist() == ["배송", "배송 속도", 2, 40, 100, 0, 0, 0]
    for frame in sheets.values():
        assert frame[["긍정 (%)", "부정 (%)", "혼합 (%)", f"{OTHER_SENTIMENT} (%)"]].sum(axis=1).eq(100).all()


@pytest.mark.parametrize("has_matched_opinion", [False, True])
def test_voc_download_excludes_unmatched_labels_and_retains_original_and_matched_classes(completed, has_matched_opinion):
    store, run_id, _ = completed
    opinions = [Issue(code_id=None, sentiment="긍정", evidence_text="오배송", missing_code="새 분류 필요")]
    if has_matched_opinion:
        opinions.append(Issue(code_id="wrong", sentiment="부정", evidence_text="오배송"))
    store.save_result(run_id, 0, "V0002", CodingResult(voc_id="V0002", response_type="opinions", issues=opinions))
    originals, issues = result_tables(store, run_id)
    before = originals.copy(deep=True), issues.copy(deep=True)
    exported = classification_frame(originals, issues)
    row = exported.loc[exported["VOC ID"].eq("V0002")]
    expected = ["배송", "오배송", "부정"] if has_matched_opinion else ["", "", ""]
    assert row[["대분류", "세부분류", "감성"]].values.tolist() == [expected]
    assert row["VOC 원문"].tolist() == ["오배송"]
    workbook = load_workbook(BytesIO(xlsx_download({"VOC별 분류": exported})))
    saved = next(row for row in list(workbook.active.values)[1:] if row[0] == "V0002")
    assert [value or "" for value in saved[2:]] == expected
    pd.testing.assert_frame_equal(originals, before[0])
    pd.testing.assert_frame_equal(issues, before[1])


def test_statistics_use_merged_chart_groups_without_changing_voc_classifications(completed):
    store, run_id, codes = completed
    originals, issues = result_tables(store, run_id)
    layout = load_code_group(initial_layout(codes), codes, ["speed", "wrong"])
    frame = statistics_frames(originals, issues, codes, layout)["세부분류"]
    assert frame.iloc[0].tolist() == ["배송", "배송 속도/오배송", 4, 80, 25, 50, 25, 0]
    assert len(frame) == 1 and len(classification_frame(originals, issues)) == 6


def test_statistics_download_tracks_graph_merge_and_undo(completed, monkeypatch):
    media = MemoryMediaFileStorage("/mock/media")
    monkeypatch.setattr(app_test, "MemoryMediaFileStorage", lambda endpoint: media)
    app, prefix = open_results(completed, monkeypatch)

    def workbook(label):
        button = next(button for button in app.get("download_button") if button.label == label)
        content = media.get_file(button.proto.url.rsplit("/", 1)[-1]).content
        return load_workbook(BytesIO(content))

    original_vocs = list(workbook("VOC별 분류 다운로드")["VOC별 분류"].values)
    assert workbook("분류 통계 다운로드")["세부분류"].max_row == 3
    chart_event(app, "code", "merge", ["speed"], ["wrong"])
    merged = workbook("분류 통계 다운로드")
    assert list(merged["세부분류"].values)[1:] == [("배송", "배송 속도/오배송", 4, 80, 25, 50, 25, 0)]
    assert list(merged["대분류"].values)[1:] == [("배송", 4, 80, 25, 50, 25, 0)]
    assert list(workbook("VOC별 분류 다운로드")["VOC별 분류"].values) == original_vocs
    app.button(key=prefix + "_undo_group").click().run()
    restored = workbook("분류 통계 다운로드")["세부분류"]
    assert [row[1] for row in list(restored.values)[1:]] == ["오배송", "배송 속도"]
    assert not app.exception and not app.error


def test_downloads_refresh_after_sentiment_edit_and_preserve_first_ai_result(completed):
    store, run_id, codes = completed
    raw = deepcopy(store.results(run_id))
    save_table_sentiment(store, store.run(run_id), "V0001", ["긍정", "중립"])
    originals, issues = result_tables(store, run_id)
    stats = statistics_frames(originals, issues, codes)
    wrong = stats["세부분류"].set_index("세부분류").loc["오배송"]
    assert wrong["부정 (%)"] == pytest.approx(200 / 3)
    assert wrong[f"{OTHER_SENTIMENT} (%)"] == pytest.approx(100 / 3)
    assert stats["대분류"].iloc[0].tolist() == ["배송", 4, 80, 50, 50, 0, 0]
    vocs = classification_frame(originals, issues)
    assert vocs.loc[vocs["VOC ID"].eq("V0001") & vocs["세부분류"].eq("오배송"), "감성"].tolist() == ["중립"]
    assert store.results(run_id) == raw


def test_xlsx_keeps_ids_text_literal_originals_numeric_percentages_and_empty_sheets(completed):
    store, run_id, codes = completed
    vocs = pd.DataFrame([["000001", "=SUM(1,2)", "배송", "오배송", "부정"]],
                       columns=["VOC ID", "VOC 원문", "대분류", "세부분류", "감성"])
    workbook = load_workbook(BytesIO(xlsx_download({"VOC별 분류": vocs})))
    assert workbook["VOC별 분류"]["A2"].value == "000001"
    assert workbook["VOC별 분류"]["A2"].data_type == "s"
    assert workbook["VOC별 분류"]["B2"].value == "=SUM(1,2)"
    assert workbook["VOC별 분류"]["B2"].data_type == "s"
    originals, issues = result_tables(store, run_id)
    workbook = load_workbook(BytesIO(xlsx_download(statistics_frames(originals, issues, codes))))
    assert workbook.sheetnames == ["대분류", "세부분류"]
    assert workbook["대분류"]["C2"].value == 80 and workbook["대분류"]["C2"].data_type == "n"
    assert workbook["대분류"]["C2"].number_format == "0.0"
    assert workbook["세부분류"].freeze_panes == "A2"
    assert workbook["세부분류"].auto_filter.ref == "A1:H3"
    empty = statistics_frames(originals, issues.iloc[:0], codes)
    workbook = load_workbook(BytesIO(xlsx_download(empty)))
    assert workbook["대분류"].max_row == workbook["세부분류"].max_row == 1


def test_ui_has_two_main_downloads_and_no_popup_download_menu(completed, monkeypatch):
    app, prefix = open_results(completed, monkeypatch)
    assert [button.label for button in app.get("download_button")] == ["VOC별 분류 다운로드", "분류 통계 다운로드"]
    assert not any(button.label == "내려받기" for button in app.button)
    assert not any(box.key == prefix + "_download_kind" for box in app.selectbox)
    chart_event(app, "code", "open", ["wrong"])
    table_event(app, "filter", column="감성", values=[])
    assert len(app.get("download_button")) == 2 and not app.exception and not app.error
