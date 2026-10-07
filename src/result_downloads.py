"""전체 분류 결과를 원문별 결과와 분류별 통계 엑셀로 내보낸다."""

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.category_summary import SENTIMENT_COLORS, response_sentiments
from src.chart_data import COUNT, PERCENT
from src.grouping import group_results
from src.result_groups import grouped_chart, initial_layout
from src.response_basis import ALL_BASIS, NO_CONTENT, basis_responses

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VOC_COLUMNS = ["VOC ID", "VOC 원문", "대분류", "세부분류", "감성"]
STAT_COLUMNS = ["전체 건수", "%", *[f"{label} (%)" for label in SENTIMENT_COLORS]]


def download_tables(kind, originals, issues, codes, book, basis=ALL_BASIS, layout=None):
    """화면 필터와 별개로 선택한 응답 기준의 전체 자료를 내보낸다."""
    view = basis_responses(originals, basis)
    opinions = issues if issues.empty else issues[issues["VOC ID"].isin(view["VOC ID"])]
    sheets = ({"VOC별 분류": classification_frame(view, opinions)} if kind == "vocs" else
              statistics_frames(view, opinions, codes, layout))
    sheets["집계 기준"] = pd.DataFrame([
        ("응답 기준", basis),
        ("전체 응답 수", originals["VOC ID"].nunique()),
        ("무응답 수", originals.loc[originals["응답 상태"].eq(NO_CONTENT), "VOC ID"].nunique()),
        ("분모 응답 수", view["VOC ID"].nunique()),
        ("유효 기준 정의", "무응답으로 분류된 응답만 제외. 중립 의견 포함."),
        ("분류 기준표", book.get("name") or ""),
        ("분류 기준표 버전", book["version"]),
        ("조회 범위", "선택한 응답 기준의 전체 자료. 화면의 감성·분류 필터는 적용하지 않음."),
        ("분류 집계", "각 분류 안에서 VOC ID 중복 제거. 복수 분류 응답은 각 분류에 포함."),
        ("분류 비율 (%)", "분류별 고유 VOC 수 / 분모 응답 수 × 100. 합계는 100%를 넘을 수 있음."),
        ("분류 내 감성 (%)", "해당 분류의 고유 VOC 수 기준. 긍정·부정 동시 응답은 혼합으로 별도 집계."),
    ], columns=["항목", "값"])
    return sheets


def classification_frame(originals, issues):
    if originals.empty:
        return pd.DataFrame(columns=VOC_COLUMNS)
    classifications = issues.reindex(columns=["VOC ID", "코드 ID", "대분류", "세부분류", "감성"])
    classifications = classifications.loc[classifications["코드 ID"].fillna("").ne("")].drop(columns="코드 ID")
    # 같은 분류·감성의 반복 근거는 한 행. 다른 분류와 다른 VOC ID는 보존한다.
    return originals.reindex(columns=["VOC ID", "VOC 원문"]).merge(
        classifications, on="VOC ID", how="left", sort=False).reindex(columns=VOC_COLUMNS).fillna("").drop_duplicates()


def statistics_frames(originals, issues, codes, layout=None):
    grouped = group_results(originals, issues, codes)
    layout = layout or initial_layout(codes)
    code_map = {code.id: code for code in codes}
    sheets = {}
    for kind, sheet_name in [("category", "대분류"), ("code", "세부분류")]:
        rows = []
        for row in grouped_chart(grouped, codes, layout, kind).to_dict("records"):
            opinions = grouped.issues[grouped.issues["코드 ID"].isin(row["code_ids"])]
            view = originals[originals["VOC ID"].isin(opinions["VOC ID"])]
            sentiments = response_sentiments(view, opinions).set_index("감성")[PERCENT]
            if kind == "category":
                labels = {"대분류": row["분류"]}
            else:
                members = [code_map[identifier] for identifier in row["members"]]
                parents = list(dict.fromkeys(code.category for code in members))
                labels = {"대분류": "/".join(parents), "세부분류":
                    "/".join(code.name for code in members) if len(parents) == 1 else row["분류"]}
            rows.append({**labels, "전체 건수": row[COUNT], "%": row[PERCENT],
                         **{f"{label} (%)": sentiments[label] for label in SENTIMENT_COLORS}})
        labels = ["대분류"] if kind == "category" else ["대분류", "세부분류"]
        sheets[sheet_name] = pd.DataFrame(rows, columns=[*labels, *STAT_COLUMNS])
    return sheets


def xlsx_download(sheets):
    """ID·원문은 텍스트로, 건수·비율은 숫자로 저장한다."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append(list(values))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for index, label in enumerate(frame.columns, 1):
            width = {"VOC ID": 14, "VOC 원문": 70, "대분류": 24, "세부분류": 32, "감성": 18,
                     "항목": 24, "값": 85}.get(label, 16)
            if "무응답" in label:
                width = 30
            sheet.column_dimensions[get_column_letter(index)].width = width
            header = sheet.cell(1, index)
            header.font = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
            header.fill = PatternFill("solid", fgColor="1E3A8A")
            header.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    # 원문이 '='로 시작해도 수식으로 실행하지 않고 그대로 보존한다.
                    cell.data_type = "s"
                    cell.number_format = "@"
                else:
                    label = frame.columns[cell.column - 1]
                    cell.number_format = "0.0" if label == "%" or "(%)" in label else "#,##0"
        sheet.sheet_view.showGridLines = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
